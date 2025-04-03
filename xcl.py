import openpyxl

wb = openpyxl.load_workbook("transactions.xlsx")
print(wb.sheetnames)

sheet = wb["Sheet1"]
column = sheet["a"]
print(column)


cell = sheet["a1"]
print(cell.value)
cell.value = 1
print(cell.value)
print(cell.row)
print(cell.column)
print(cell.coordinate)

cell = sheet.cell(row = 1, column = 1)
print(sheet.max_row)
print(sheet.max_column)

sheet.append([1,2,3])
wb.save("transactions2.xlsx")